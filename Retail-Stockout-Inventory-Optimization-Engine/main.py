from pathlib import Path
import sqlite3
import warnings

import numpy as np
import pandas as pd
from sklearn.ensemble import RandomForestRegressor

warnings.filterwarnings("ignore", category=UserWarning)

ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
OUTPUT_DIR = ROOT / "output"
SALES_PATH = DATA_DIR / "sales_history.csv"
INVENTORY_PATH = DATA_DIR / "inventory.csv"
REPORT_PATH = OUTPUT_DIR / "inventory_optimization_report.xlsx"
SUMMARY_PATH = OUTPUT_DIR / "sql_product_summary.csv"

FORECAST_HORIZON_DAYS = 7
SERVICE_LEVEL_Z = 1.65  # Approximate 95% one-sided service level


def create_sample_data():
    """Create reproducible synthetic sales and inventory CSVs."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(42)
    today = pd.Timestamp.today().normalize()
    dates = pd.date_range(end=today - pd.Timedelta(days=1), periods=180, freq="D")

    products = [
        ("P001", "Wireless Mouse", "Electronics", 18, 4, 650.0, 14),
        ("P002", "USB-C Cable", "Electronics", 55, 3, 180.0, 14),
        ("P003", "Water Bottle", "Lifestyle", 12, 6, 250.0, 21),
        ("P004", "Notebook Pack", "Stationery", 70, 2, 90.0, 14),
        ("P005", "Desk Lamp", "Home", 5, 8, 950.0, 21),
        ("P006", "Backpack", "Lifestyle", 24, 5, 1200.0, 14),
        ("P007", "Bluetooth Speaker", "Electronics", 9, 7, 1750.0, 21),
        ("P008", "Pen Set", "Stationery", 110, 2, 60.0, 14),
        ("P009", "Coffee Mug", "Home", 42, 4, 220.0, 14),
        ("P010", "Keyboard", "Electronics", 7, 6, 1450.0, 21),
        ("P011", "Phone Stand", "Accessories", 35, 3, 300.0, 14),
        ("P012", "Lunch Box", "Lifestyle", 16, 5, 400.0, 14),
    ]

    sales_rows = []
    inventory_rows = []
    base_demands = {
        "P001": 5.5, "P002": 12.0, "P003": 4.0, "P004": 10.0,
        "P005": 2.4, "P006": 3.5, "P007": 2.0, "P008": 16.0,
        "P009": 5.0, "P010": 2.8, "P011": 4.5, "P012": 3.8,
    }
    for pid, name, category, stock, lead, cost, cover in products:
        base = base_demands[pid]
        for d in dates:
            weekday_factor = 1.18 if d.dayofweek in (4, 5) else 0.92 if d.dayofweek == 0 else 1.0
            seasonal_factor = 1 + 0.12 * np.sin(2 * np.pi * d.dayofyear / 365.25)
            noise = rng.normal(1.0, 0.22)
            units = max(0, int(round(base * weekday_factor * seasonal_factor * noise)))
            sales_rows.append((d.date().isoformat(), pid, units))
        inventory_rows.append((pid, name, category, stock, lead, cost, cover))

    pd.DataFrame(sales_rows, columns=["date", "product_id", "units_sold"]).to_csv(SALES_PATH, index=False)
    pd.DataFrame(
        inventory_rows,
        columns=["product_id", "product_name", "category", "current_stock",
                 "supplier_lead_time_days", "unit_cost", "target_cover_days"]
    ).to_csv(INVENTORY_PATH, index=False)
    print("Created synthetic sample CSV files in data/.")


def load_and_validate_data():
    if not SALES_PATH.exists() or not INVENTORY_PATH.exists():
        create_sample_data()

    sales = pd.read_csv(SALES_PATH)
    inventory = pd.read_csv(INVENTORY_PATH)

    required_sales = {"date", "product_id", "units_sold"}
    required_inventory = {
        "product_id", "product_name", "category", "current_stock",
        "supplier_lead_time_days", "unit_cost", "target_cover_days"
    }
    if not required_sales.issubset(sales.columns):
        raise ValueError(f"sales_history.csv must include columns: {sorted(required_sales)}")
    if not required_inventory.issubset(inventory.columns):
        raise ValueError(f"inventory.csv must include columns: {sorted(required_inventory)}")

    sales["date"] = pd.to_datetime(sales["date"], errors="coerce")
    sales["product_id"] = sales["product_id"].astype(str).str.strip().str.upper()
    inventory["product_id"] = inventory["product_id"].astype(str).str.strip().str.upper()
    sales["units_sold"] = pd.to_numeric(sales["units_sold"], errors="coerce").fillna(0).clip(lower=0)
    for col in ["current_stock", "supplier_lead_time_days", "unit_cost", "target_cover_days"]:
        inventory[col] = pd.to_numeric(inventory[col], errors="coerce")
    sales = sales.dropna(subset=["date", "product_id"])
    inventory = inventory.dropna(subset=["product_id", "current_stock", "supplier_lead_time_days"])
    inventory["current_stock"] = inventory["current_stock"].clip(lower=0)
    inventory["supplier_lead_time_days"] = inventory["supplier_lead_time_days"].clip(lower=1)
    inventory["unit_cost"] = inventory["unit_cost"].fillna(0).clip(lower=0)
    inventory["target_cover_days"] = inventory["target_cover_days"].fillna(14).clip(lower=1)

    if inventory["product_id"].duplicated().any():
        raise ValueError("inventory.csv must contain only one row per product_id.")
    if sales.empty or inventory.empty:
        raise ValueError("Input data is empty after validation.")
    return sales, inventory


def forecast_product_demand(product_sales, horizon=FORECAST_HORIZON_DAYS):
    """Forecast future daily units for one product using a small Random Forest model."""
    daily = (
        product_sales.groupby("date", as_index=False)["units_sold"]
        .sum()
        .sort_values("date")
    )
    # Fill missing dates with zero sales to make the time series regular.
    idx = pd.date_range(daily["date"].min(), daily["date"].max(), freq="D")
    daily = daily.set_index("date").reindex(idx, fill_value=0).rename_axis("date").reset_index()
    daily["units_sold"] = daily["units_sold"].astype(float)

    if len(daily) < 21:
        average = float(daily["units_sold"].mean()) if len(daily) else 0.0
        return [max(0.0, average)] * horizon, float(daily["units_sold"].std(ddof=1) or 0.0)

    frame = daily.copy()
    frame["day_of_week"] = frame["date"].dt.dayofweek
    frame["day_of_year"] = frame["date"].dt.dayofyear
    frame["lag_1"] = frame["units_sold"].shift(1)
    frame["lag_7"] = frame["units_sold"].shift(7)
    frame["rolling_mean_7"] = frame["units_sold"].shift(1).rolling(7).mean()
    frame = frame.dropna().reset_index(drop=True)

    feature_cols = ["day_of_week", "day_of_year", "lag_1", "lag_7", "rolling_mean_7"]
    X = frame[feature_cols]
    y = frame["units_sold"]
    model = RandomForestRegressor(
        n_estimators=120, max_depth=8, min_samples_leaf=2, random_state=42, n_jobs=-1
    )
    model.fit(X, y)

    history = daily.set_index("date")["units_sold"].to_dict()
    last_date = daily["date"].max()
    predictions = []
    for step in range(1, horizon + 1):
        future_date = last_date + pd.Timedelta(days=step)
        lag1_date = future_date - pd.Timedelta(days=1)
        lag7_date = future_date - pd.Timedelta(days=7)
        lag1 = float(history.get(lag1_date, daily["units_sold"].iloc[-1]))
        lag7 = float(history.get(lag7_date, daily["units_sold"].mean()))
        recent = [
            float(history.get(future_date - pd.Timedelta(days=i), daily["units_sold"].mean()))
            for i in range(1, 8)
        ]
        row = pd.DataFrame([{
            "day_of_week": future_date.dayofweek,
            "day_of_year": future_date.dayofyear,
            "lag_1": lag1,
            "lag_7": lag7,
            "rolling_mean_7": float(np.mean(recent)),
        }])
        prediction = max(0.0, float(model.predict(row[feature_cols])[0]))
        predictions.append(prediction)
        history[future_date] = prediction

    demand_std = float(daily["units_sold"].std(ddof=1))
    if not np.isfinite(demand_std):
        demand_std = 0.0
    return predictions, demand_std


def build_recommendations(sales, inventory):
    rows = []
    for _, item in inventory.iterrows():
        product_sales = sales[sales["product_id"] == item["product_id"]]
        predictions, demand_std = forecast_product_demand(product_sales)
        forecast_7d = float(np.sum(predictions))
        forecast_daily = forecast_7d / FORECAST_HORIZON_DAYS
        lead_time = int(max(1, round(item["supplier_lead_time_days"])))
        current_stock = float(item["current_stock"])

        # Safety stock uses historical demand variability and supplier lead time.
        safety_stock = SERVICE_LEVEL_Z * demand_std * np.sqrt(lead_time)
        lead_time_demand = forecast_daily * lead_time
        reorder_point = lead_time_demand + safety_stock
        target_stock = forecast_daily * (lead_time + float(item["target_cover_days"])) + safety_stock
        suggested_order = max(0, int(np.ceil(target_stock - current_stock)))

        if current_stock <= lead_time_demand:
            risk = "Critical"
            action = "Urgent replenishment; check supplier availability"
        elif current_stock <= reorder_point:
            risk = "At Risk"
            action = "Place replenishment order"
        elif current_stock > target_stock * 1.5 and target_stock > 0:
            risk = "Overstocked"
            action = "Pause/reduce orders; review slow-moving stock"
        else:
            risk = "Healthy"
            action = "Monitor demand and stock levels"

        rows.append({
            "product_id": item["product_id"],
            "product_name": item["product_name"],
            "category": item["category"],
            "current_stock": round(current_stock, 2),
            "supplier_lead_time_days": lead_time,
            "avg_daily_demand_history": round(float(product_sales["units_sold"].mean()), 2),
            "forecast_daily_demand": round(forecast_daily, 2),
            "forecast_next_7_days": round(forecast_7d, 2),
            "historical_demand_std": round(demand_std, 2),
            "safety_stock": round(float(safety_stock), 2),
            "lead_time_demand": round(float(lead_time_demand), 2),
            "reorder_point": round(float(reorder_point), 2),
            "target_stock_level": round(float(target_stock), 2),
            "suggested_order_quantity": suggested_order,
            "inventory_risk": risk,
            "recommended_action": action,
            "unit_cost": round(float(item["unit_cost"]), 2),
            "estimated_order_value": round(suggested_order * float(item["unit_cost"]), 2),
        })
    return pd.DataFrame(rows).sort_values(
        by=["inventory_risk", "suggested_order_quantity"],
        ascending=[True, False]
    ).reset_index(drop=True)


def create_sql_summary(sales, inventory):
    """Run a genuine SQL aggregation over the validated CSV data using SQLite."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(":memory:") as conn:
        sales.to_sql("sales", conn, index=False, if_exists="replace")
        inventory.to_sql("inventory", conn, index=False, if_exists="replace")
        query = """
        SELECT
            i.product_id,
            i.product_name,
            i.category,
            i.current_stock,
            ROUND(AVG(s.units_sold), 2) AS avg_daily_units_sold,
            SUM(s.units_sold) AS total_units_sold,
            COUNT(DISTINCT s.date) AS sales_days,
            ROUND(i.current_stock * i.unit_cost, 2) AS current_inventory_value
        FROM inventory i
        LEFT JOIN sales s ON i.product_id = s.product_id
        GROUP BY i.product_id, i.product_name, i.category, i.current_stock, i.unit_cost
        ORDER BY total_units_sold DESC
        """
        summary = pd.read_sql_query(query, conn)
    summary.to_csv(SUMMARY_PATH, index=False)
    return summary


def export_excel(recommendations, sql_summary, sales):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    risk_counts = recommendations["inventory_risk"].value_counts()
    kpis = pd.DataFrame([
        ["Total products", int(len(recommendations))],
        ["Critical products", int(risk_counts.get("Critical", 0))],
        ["At-risk products", int(risk_counts.get("At Risk", 0))],
        ["Healthy products", int(risk_counts.get("Healthy", 0))],
        ["Overstocked products", int(risk_counts.get("Overstocked", 0))],
        ["Total suggested units to order", int(recommendations["suggested_order_quantity"].sum())],
        ["Estimated replenishment value", round(float(recommendations["estimated_order_value"].sum()), 2)],
        ["Current inventory value", round(float((recommendations["current_stock"] * recommendations["unit_cost"]).sum()), 2)],
    ], columns=["KPI", "Value"])

    with pd.ExcelWriter(REPORT_PATH, engine="openpyxl") as writer:
        kpis.to_excel(writer, sheet_name="Dashboard", index=False, startrow=0)
        recommendations.to_excel(writer, sheet_name="Recommendations", index=False)
        sql_summary.to_excel(writer, sheet_name="SQL Summary", index=False)
        sales.sort_values(["date", "product_id"]).to_excel(writer, sheet_name="Sales History", index=False)

        # Basic dashboard formatting and useful filters.
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for col_idx, column_cells in enumerate(ws.columns, start=1):
                values = [str(c.value) if c.value is not None else "" for c in list(column_cells)[:200]]
                width = min(max(max((len(v) for v in values), default=10) + 2, 12), 32)
                ws.column_dimensions[get_column_letter(col_idx)].width = width
            ws.row_dimensions[1].height = 28

        # Highlight the dashboard KPI header and add a simple chart.
        dashboard = writer.book["Dashboard"]
        dashboard["A1"].font = Font(bold=True, color="FFFFFF", size=12)
        dashboard["B1"].font = Font(bold=True, color="FFFFFF", size=12)
        dashboard.freeze_panes = "A2"

        from openpyxl.chart import BarChart, Reference
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Inventory Risk Counts"
        chart.y_axis.title = "Risk category"
        chart.x_axis.title = "Products"
        risk_order = ["Critical", "At Risk", "Healthy", "Overstocked"]
        start_row = len(kpis) + 3
        dashboard.cell(start_row, 1, "Risk Category")
        dashboard.cell(start_row, 2, "Product Count")
        for offset, risk in enumerate(risk_order, start=1):
            dashboard.cell(start_row + offset, 1, risk)
            dashboard.cell(start_row + offset, 2, int(risk_counts.get(risk, 0)))
        for cell in dashboard[start_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        data = Reference(dashboard, min_col=2, min_row=start_row, max_row=start_row + len(risk_order))
        cats = Reference(dashboard, min_col=1, min_row=start_row + 1, max_row=start_row + len(risk_order))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 7
        chart.width = 14
        dashboard.add_chart(chart, "D2")

    print(f"Excel report created: {REPORT_PATH}")


def main():
    sales, inventory = load_and_validate_data()
    recommendations = build_recommendations(sales, inventory)
    sql_summary = create_sql_summary(sales, inventory)
    export_excel(recommendations, sql_summary, sales)

    display_cols = [
        "product_id", "product_name", "current_stock", "forecast_daily_demand",
        "reorder_point", "suggested_order_quantity", "inventory_risk"
    ]
    print("\nInventory recommendations:")
    print(recommendations[display_cols].to_string(index=False))
    print(f"\nSQL summary CSV: {SUMMARY_PATH}")
    print(f"Full Excel report: {REPORT_PATH}")


if __name__ == "__main__":
    main()
